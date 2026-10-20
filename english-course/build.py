#!/usr/bin/env python3
"""Build the course PDFs (plus editable Markdown / DOCX) from plain-text lesson files.

    python3 build.py              # roadmap + every book
    python3 build.py check        # validate lesson files only
    python3 build.py syllabus     # 180-day roadmap PDF
    python3 build.py book1        # Book 1: PDF + Markdown + DOCX

Lesson files live in content/bookN/dayNN.txt; see content/FORMAT.md.
"""
import html
import os
import re
import shutil
import subprocess
import sys
import urllib.request
from pathlib import Path

import jinja2
import yaml

ROOT = Path(__file__).resolve().parent
CONTENT = ROOT / "content"
TEMPLATES = ROOT / "templates"
BUILD = ROOT / "build"
OUTPUT = ROOT / "output"
FONTS = ROOT / "fonts"

FONT_URLS = {
    "NotoSans-Regular.ttf": "https://fonts.gstatic.com/s/notosans/v42/o-0mIpQlx3QUlC5A4PNB6Ryti20_6n1iPHjcz6L1SoM-jCpoiyD9A99d.ttf",
    "NotoSans-Bold.ttf": "https://fonts.gstatic.com/s/notosans/v42/o-0mIpQlx3QUlC5A4PNB6Ryti20_6n1iPHjcz6L1SoM-jCpoiyAaBN9d.ttf",
    "NotoSansSC-Regular.ttf": "https://fonts.gstatic.com/s/notosanssc/v41/k3kCo84MPvpLmixcA63oeAL7Iqp5IZJF9bmaG9_FnYw.ttf",
    "NotoSansSC-Bold.ttf": "https://fonts.gstatic.com/s/notosanssc/v41/k3kCo84MPvpLmixcA63oeAL7Iqp5IZJF9bmaGzjCnYw.ttf",
    "NotoSerif-Regular.ttf": "https://fonts.gstatic.com/s/notoserif/v33/ga6iaw1J5X9T9RW6j9bNVls-hfgvz8JcMofYTa32J4wsL2JAlAhZqFCjwA.ttf",
    "NotoSerif-Bold.ttf": "https://fonts.gstatic.com/s/notoserif/v33/ga6iaw1J5X9T9RW6j9bNVls-hfgvz8JcMofYTa32J4wsL2JAlAhZT1ejwA.ttf",
    "NotoSerif-Italic.ttf": "https://fonts.gstatic.com/s/notoserif/v33/ga6saw1J5X9T9RW6j9bNfFIMZhhWnFTyNZIQD1-_FXP0RgnaOg9MYBNLg8cP.ttf",
}

# ---------------------------------------------------------------------------
# The 120-minute plans. Each step: (name, minutes, concrete instruction).
# Instructions point at the lettered sections of the same lesson.
# ---------------------------------------------------------------------------
REGULAR_PLAN = [
    ("旧知识情境复习", 10, "遮住答案，看「热身」里的每句中文，5 秒内用英文说出来。说不出的打 ★，翻回原课把那句读 3 遍，再说一次。"),
    ("中文意图 → 英语表达", 15, "先读【B】，用中文说出这个场景里你要表达的 3 个意思；再学【E】，每组对比例句大声读 3 遍，完成「小任务」后翻到【答案】核对。"),
    ("微故事与词汇", 25, "【C】第 1 遍只看英文猜意思；第 2 遍逐句对照中文；第 3 遍配合【D】注释弄懂每个重点词；最后遮住中文，把全文大声读 2 遍。"),
    ("语音训练", 15, "按【G】：每个重点词拍着重音读 5 遍；每个句子按「/」停顿、「‿」连读读 3 遍；录下来和手机朗读版对比，找出 1 处不同。"),
    ("听力训练", 15, "用手机朗读功能播放【C】，按「四步听法」：①看英文听 ②只看中文听 ③什么都不看听 ④听写【G】里的 2 个句子，再对照原文。"),
    ("开口训练", 20, "【H】三层练习全部出声：第 1 层模仿 3 遍；第 2 层每题至少说 2 个替换版本；第 3 层只看中文情境即兴说，说完再看示范。"),
    ("输出训练", 12, "完成【H】最后的「今日输出」：录一段 1 分钟语音，或写 5–8 句话，按日期存档。"),
    ("当天复盘", 8, "做【I】检测题并对答案；在「完成清单」打勾；把「明日复习」和「七日后复习」抄进笔记本或手机日历。"),
]
REVIEW_PLAN = [
    ("旧知识情境复习", 10, "遮住答案，用英文说出「热身」的中文提示。这些句子来自本周不同的日子，说不出的打 ★。"),
    ("新情境阅读", 20, "【C】是一个用本周旧词写成的新故事：先遮住中文读 2 遍，看到虚线词（复现词）就停下来回忆意思和第几天学的，再逐句对照中文。"),
    ("听懂短对话", 15, "用手机朗读功能播放【C】里的对话部分，按「四步听法」听 4 遍，最后不看文字回答【I】的前 2 题。"),
    ("中文意思 → 英文组织", 20, "【H】第 2、3 层：只看中文情境，用本周学过的表达组织英文，先说再写，最后对照示范。"),
    ("不看原文复述", 15, "合上书，按【H】第 1 层的提示用 6–8 句英文复述【C】的故事，并录音。"),
    ("易混表达检查", 15, "做【E】的易混表达检查：每组先自己选、说出理由，再对答案；错的那组把正确例句读 5 遍。"),
    ("输出与自评", 15, "完成【H】的「今日输出」录音，再做【I】剩下的检测题。"),
    ("下阶段改善", 10, "读本课最后的「下阶段改善建议」，从薄弱项里选 2 个，写进下周每天的热身时间里。"),
]

SCALARS = {
    "day", "week", "review", "title_cn", "title_en", "topic", "think_short",
    "goal_scene", "goal_expr", "goal_vocab", "scene", "story_title", "story_type",
    "mimic", "output", "output_sample", "res", "warmup_hint",
}
LISTS = {
    "done": "done", "pnote": "pron_notes", "lfocus": "listen_focus", "dict": "dictation",
    "next1": "review_next_day", "next7": "review_next_week", "improve": "improve",
}
SPLITS = {
    # key: (destination, field names, minimum number of fields)
    "warmup": ("warmup", ["cn", "en", "src"], 2),
    "sp": ("scene_points", ["label", "text"], 2),
    "core": ("core", ["expr", "use", "swap", "tone"], 4),
    "pw": ("pron_words", ["word", "ipa", "stress", "tip"], 4),
    "ps": ("pron_sents", ["sent", "note"], 2),
    "swap": ("swaps", ["base", "hint", "sample"], 3),
    "free": ("frees", ["cn", "sample"], 2),
    "check": ("checks", ["q", "a"], 2),
    "assess": ("assess", ["item", "days"], 2),
    "weak": ("weak", ["item", "fix"], 2),
    "plan": ("plan", ["name", "min", "how"], 3),
}
VOCAB_FIELDS = ["form", "head", "pos", "ipa", "meaning", "scene", "colloc", "ex_en", "ex_cn", "pitfall"]
DIFF_FIELDS = {"d_cn": "cn", "d_en": "en", "d_why": "why", "d_when": "when", "d_vs": "vs"}
REQUIRED = ["day", "week", "title_cn", "title_en", "topic", "goal_scene", "goal_expr", "goal_vocab",
            "scene", "story_title", "mimic", "output", "output_sample"]

IRREGULAR = {
    "be": "am is are was were been being", "have": "has had having", "do": "does did done doing",
    "go": "goes went gone going", "get": "gets got gotten getting", "make": "makes made making",
    "take": "takes took taken taking", "come": "comes came coming", "see": "sees saw seen seeing",
    "say": "says said saying", "tell": "tells told telling", "speak": "speaks spoke spoken speaking",
    "eat": "eats ate eaten eating", "drink": "drinks drank drunk drinking", "buy": "buys bought buying",
    "think": "thinks thought thinking", "bring": "brings brought bringing", "leave": "leaves left leaving",
    "meet": "meets met meeting", "send": "sends sent sending", "spend": "spends spent spending",
    "feel": "feels felt feeling", "find": "finds found finding", "give": "gives gave given giving",
    "know": "knows knew known knowing", "write": "writes wrote written writing", "begin": "begins began begun",
    "break": "breaks broke broken breaking", "forget": "forgets forgot forgotten forgetting",
    "lose": "loses lost losing", "pay": "pays paid paying", "sit": "sits sat sitting", "stand": "stands stood",
    "understand": "understands understood understanding", "wake": "wakes woke woken waking",
    "catch": "catches caught catching", "teach": "teaches taught teaching", "sleep": "sleeps slept sleeping",
    "run": "runs ran running", "choose": "chooses chose chosen choosing", "fall": "falls fell fallen falling",
    "hold": "holds held holding", "keep": "keeps kept keeping", "put": "puts putting", "set": "sets setting",
    "let": "lets letting", "hurt": "hurts hurting", "cost": "costs costing", "read": "reads reading",
    "hear": "hears heard hearing", "lead": "leads led leading", "win": "wins won winning",
    "drive": "drives drove driven driving", "ride": "rides rode ridden riding", "fly": "flies flew flown flying",
    "grow": "grows grew grown growing", "show": "shows showed shown showing", "wear": "wears wore worn wearing",
}


class LessonError(Exception):
    pass


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
def split_fields(val, names, minimum, where):
    parts = [p.strip() for p in val.split(" | ")]
    if not minimum <= len(parts) <= len(names):
        raise LessonError(f"{where}: expected {minimum}-{len(names)} fields separated by ' | ', got {len(parts)}: {val[:70]}")
    parts += [""] * (len(names) - len(parts))
    return dict(zip(names, parts))


def parse_day(path):
    d = {dest: [] for dest, _, _ in SPLITS.values()}
    d.update({dest: [] for dest in LISTS.values()})
    d.update(sentences=[], paragraph_breaks=[], vocab=[], diffs=[], review=False)
    for n, raw in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        where = f"{path.name}:{n}"
        if not line.startswith("@") or ":" not in line:
            raise LessonError(f"{where}: line must look like '@key: value'")
        key, _, val = line[1:].partition(":")
        key, val = key.strip(), val.strip()
        if key in SCALARS:
            d[key] = val
        elif key in LISTS:
            d[LISTS[key]].append(val)
        elif key in SPLITS:
            dest, names, minimum = SPLITS[key]
            d[dest].append(split_fields(val, names, minimum, where))
        elif key == "s":
            s = split_fields(val, ["en", "cn"], 2, where)
            m = re.match(r"^([A-Z][a-zA-Z]+(?: [A-Z][a-z]+)?): (.*)$", s["en"])
            s["speaker"], s["text"] = (m.group(1), m.group(2)) if m else ("", s["en"])
            s["n"] = len(d["sentences"]) + 1
            d["sentences"].append(s)
        elif key == "br":
            d["paragraph_breaks"].append(len(d["sentences"]))
        elif key == "v":
            v = split_fields(val, VOCAB_FIELDS, len(VOCAB_FIELDS), where)
            if v["head"] in ("", "="):
                v["head"] = v["form"]
            v["src"] = where
            d["vocab"].append(v)
        elif key == "diff":
            d["diffs"].append({"title": val, "task": "", "task_answer": ""})
        elif key in DIFF_FIELDS:
            if not d["diffs"]:
                raise LessonError(f"{where}: @{key} before any @diff")
            d["diffs"][-1][DIFF_FIELDS[key]] = val
        elif key == "d_task":
            t = split_fields(val, ["q", "a"], 2, where)
            d["diffs"][-1].update(task=t["q"], task_answer=t["a"])
        else:
            raise LessonError(f"{where}: unknown key @{key}")
    missing = [k for k in REQUIRED if not d.get(k)]
    if missing:
        raise LessonError(f"{path.name}: missing {missing}")
    d["day"], d["week"] = int(d["day"]), int(d["week"])
    d["review"] = str(d["review"]).lower() in ("1", "true", "yes")
    if d["plan"]:
        d["plan"] = [(p["name"], int(p["min"]), p["how"]) for p in d["plan"]]
    else:
        d["plan"] = REVIEW_PLAN if d["review"] else REGULAR_PLAN
    locate_vocab(d, path.name)
    d["word_count"] = sum(1 for v in d["vocab"] if v["pos"] != "phr.")
    d["phrase_count"] = sum(1 for v in d["vocab"] if v["pos"] == "phr.")
    d["story_words"] = len(re.findall(r"[A-Za-z][A-Za-z'’-]*", plain(" ".join(s["text"] for s in d["sentences"]))))
    return d


def locate_vocab(d, fname):
    """Number the vocab in order of appearance and record which sentence it is in."""
    by_form = {v["form"].lower(): v for v in d["vocab"]}
    d["unmatched_bold"] = []
    order = 0
    for s in d["sentences"]:
        for b in re.findall(r"\*\*(.+?)\*\*", s["text"]):
            v = by_form.get(b.lower())
            if v is None:
                d["unmatched_bold"].append(f"{fname} S{s['n']}: **{b}**")
            elif "sentence" not in v:
                order += 1
                v.update(sentence=s["n"], num=order)
    for v in d["vocab"]:
        if "sentence" not in v:
            raise LessonError(f"{v['src']}: '{v['form']}' is not bolded in any story sentence")
    d["vocab"].sort(key=lambda v: v["num"])


# ---------------------------------------------------------------------------
# Vocabulary tracking across days
# ---------------------------------------------------------------------------
def word_variants(word):
    w = word.lower().strip()
    w = re.sub(r"\b(sb|sth|one's)\b", "", w)
    w = re.sub(r"\s+", " ", w).strip()
    parts = w.split(" ")
    head, rest = parts[0], " ".join(parts[1:])
    forms = {head}
    if head in IRREGULAR:
        forms.update(IRREGULAR[head].split())
    forms.update({head + "s", head + "es", head + "ed", head + "d", head + "ing"})
    if head.endswith("y"):
        forms.update({head[:-1] + "ies", head[:-1] + "ied"})
    if head.endswith("e"):
        forms.add(head[:-1] + "ing")
    if re.search(r"[^aeiou][aeiou][bdglmnprt]$", head):
        forms.update({head + head[-1] + "ed", head + head[-1] + "ing"})
    return {(f + " " + rest).strip() for f in forms}


def lesson_english(d):
    """All English the learner meets in a lesson (story, examples, answers)."""
    chunks = [s["text"] for s in d["sentences"]]
    chunks += [w["en"] for w in d["warmup"]] + [c["expr"] for c in d["core"]]
    chunks += [x["sample"] for x in d["swaps"]] + [x["sample"] for x in d["frees"]]
    chunks += [d.get("output_sample", "")] + [x.get("en", "") + " " + x.get("vs", "") for x in d["diffs"]]
    return " ".join(plain(c) for c in chunks).lower()


def track_vocab(days):
    index = []
    texts = {d["day"]: lesson_english(d) for d in days}
    for d in days:
        for v in d["vocab"]:
            pats = [re.compile(r"(?<![a-z'])" + re.escape(f) + r"(?![a-z])") for f in word_variants(v["head"]) | {v["form"].lower()}]
            later = [day for day, t in texts.items() if day > d["day"] and any(p.search(t) for p in pats)]
            v["recur"] = later
            v["day"] = d["day"]
            index.append(v)
    return index


def validate(days):
    problems, warnings = [], []
    seen = {}
    for d in days:
        tag = f"day{d['day']:02d}"
        total = sum(m for _, m, _ in d["plan"])
        if total != 120:
            problems.append(f"{tag}: plan adds up to {total} minutes, not 120")
        for v in d["vocab"]:
            key = re.sub(r"\s+", " ", re.sub(r"\(.*?\)", "", v["head"].lower())).strip()
            if key in seen:
                problems.append(f"duplicate headword '{v['head']}' in {v['src']} (first in {seen[key]})")
            seen.setdefault(key, v["src"])
        warnings += [f"bold word without a note: {u}" for u in d["unmatched_bold"]]
        if not d["review"]:
            if not 12 <= d["word_count"] <= 22:
                warnings.append(f"{tag}: {d['word_count']} key words (aim 15-20)")
            if not 3 <= d["phrase_count"] <= 6:
                warnings.append(f"{tag}: {d['phrase_count']} phrases (aim 3-5)")
            if not 2 <= len(d["diffs"]) <= 4:
                warnings.append(f"{tag}: {len(d['diffs'])} Chinese/English differences (aim 2-4)")
            for x in d["diffs"]:
                for f in ("cn", "en", "why", "when", "vs", "task"):
                    if not x.get(f):
                        problems.append(f"{tag}: difference '{x['title']}' is missing '{f}'")
        if not 5 <= len(d["core"]) <= 8:
            warnings.append(f"{tag}: {len(d['core'])} core expressions (aim 5-8)")
        if len(d["swaps"]) < 3 or len(d["frees"]) < 3:
            warnings.append(f"{tag}: speaking practice needs 3+ substitution and 3+ free items")
        if len(d["checks"]) < 4:
            warnings.append(f"{tag}: only {len(d['checks'])} check questions")
        if not d["review_next_day"] or not d["review_next_week"]:
            problems.append(f"{tag}: missing next-day / seven-day review lists")
        for s in d["sentences"]:
            for r in re.findall(r"\[\[(.+?)\]\]", s["text"]):
                if first_day(r) is None:
                    warnings.append(f"{tag} S{s['n']}: recycled word [[{r}]] not found in earlier lessons")
    return problems, warnings


FORM_DAY = {}  # lower-cased form / headword -> day first taught (filled in load_book)


# ---------------------------------------------------------------------------
# Inline markup
# ---------------------------------------------------------------------------
def md(text):
    """**bold**, `ipa`, [[recycled]], ‿ linking mark, \\n line break."""
    if text is None:
        return ""
    s = html.escape(str(text), quote=False)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"`(.+?)`", r'<span class="ipa">\1</span>', s)
    s = re.sub(r"\[\[(.+?)\]\]", lambda m: recycled_html(m.group(1)), s)
    s = s.replace("‿", '<span class="lk">‿</span>')
    s = s.replace("\\n", "<br>")
    return jinja2.utils.markupsafe.Markup(s)


def first_day(word):
    """Day a recycled word was first taught, matching inflected forms too."""
    w = word.lower()
    if w in FORM_DAY:
        return FORM_DAY[w]
    days = [day for head, day in FORM_DAY.items() if w in word_variants(head)]
    return min(days) if days else None


def recycled_html(word):
    day = first_day(html.unescape(word))
    sup = f'<sup class="rd">D{day}</sup>' if day else ""
    return f'<span class="rcy">{word}</span>{sup}'


def story_md(text, vocab):
    """Like md(), but bold words get the number of their note in section D."""
    nums = {v["form"].lower(): v["num"] for v in vocab}
    s = str(md(text))
    s = re.sub(r"<b>(.+?)</b>", lambda m: f'<b>{m.group(1)}</b><sup class="vn">{nums.get(html.unescape(m.group(1)).lower(), "")}</sup>', s)
    return jinja2.utils.markupsafe.Markup(s)


def plain(text):
    return re.sub(r"\*\*|`|\[\[|\]\]|‿", "", str(text)).replace("\\n", " ")


def mdtext(text):
    """Inline markup for the Markdown/DOCX export."""
    s = str(text or "")
    s = re.sub(r"\[\[(.+?)\]\]", lambda m: f"_{m.group(1)}_" + (f"(D{first_day(m.group(1))})" if first_day(m.group(1)) else ""), s)
    s = s.replace("`", "").replace("\\n", " / ").replace("|", "／")
    return s


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------
def ensure_fonts():
    FONTS.mkdir(exist_ok=True)
    for name, url in FONT_URLS.items():
        if not (FONTS / name).exists():
            print("downloading", name)
            urllib.request.urlretrieve(url, FONTS / name)


def env(autoescape=True):
    e = jinja2.Environment(loader=jinja2.FileSystemLoader(TEMPLATES), autoescape=autoescape,
                           trim_blocks=True, lstrip_blocks=True)
    e.filters.update(md=md, story_md=story_md, plain=plain, mdtext=mdtext)
    return e


def chrome_pdf(html_path, pdf_path, attempts=3):
    """Render with Chromium. Large documents are occasionally cut short, so every template ends
    with a "THE END" marker and the render is retried until that marker is on the last page."""
    import pypdfium2 as pdfium
    npm_root = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    for attempt in range(1, attempts + 1):
        subprocess.run(["node", str(ROOT / "tools" / "render.cjs"), str(html_path), str(pdf_path)],
                       check=True, env={**os.environ, "NODE_PATH": npm_root})
        doc = pdfium.PdfDocument(str(pdf_path))
        last = doc[len(doc) - 1].get_textpage().get_text_range()
        doc.close()
        if "THE END" in last:
            return
        print(f"render of {html_path.name} was cut short ({attempt}/{attempts}), retrying")
    raise RuntimeError(f"{html_path.name}: Chromium kept producing an incomplete PDF")


def find_pages(pdf_path, patterns):
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(pdf_path))
    found = {}
    for i in range(len(doc)):
        text = doc[i].get_textpage().get_text_range()
        for label, rx in patterns:
            if label not in found and re.search(rx, text):
                found[label] = i + 1
        for m in re.finditer(r"D\s*A\s*Y\s*(\d{2,3})\b", text):
            found.setdefault(f"DAY {int(m.group(1))}", i + 1)
    return found, len(doc)


def build_pdf(template, context, out_name, footer, patterns=()):
    BUILD.mkdir(exist_ok=True)
    OUTPUT.mkdir(exist_ok=True)
    shutil.copy(TEMPLATES / "style.css", BUILD / "style.css")
    tpl = env().get_template(template)
    stem = Path(out_name).stem
    html_path, raw = BUILD / f"{stem}.html", BUILD / f"{stem}.pdf"
    pages = {}
    for _ in range(2):  # pass 1 finds page numbers, pass 2 prints them in the contents
        html_path.write_text(tpl.render(pages=pages, footer=footer, **context), encoding="utf-8")
        chrome_pdf(html_path, raw)
        pages, count = find_pages(raw, patterns)
    shutil.copy(raw, OUTPUT / out_name)
    print(f"wrote output/{out_name} ({count} pages)")


def build_editable(context, stem):
    """Markdown source and a DOCX converted from it with pandoc."""
    text = env(autoescape=False).get_template("book.md.j2").render(**context)
    md_path = OUTPUT / f"{stem}.md"
    md_path.write_text(text, encoding="utf-8")
    print(f"wrote output/{md_path.name}")
    if shutil.which("pandoc"):
        ref = BUILD / "reference.docx"
        make_reference_docx(ref)
        docx = OUTPUT / f"{stem}.docx"
        subprocess.run(["pandoc", str(md_path), "-f", "markdown+pipe_tables", "-o", str(docx),
                        "--reference-doc", str(ref), "--toc", "--toc-depth=2"], check=True)
        print(f"wrote output/{docx.name}")
    else:
        print("pandoc not found: skipped DOCX")


def make_reference_docx(path):
    """pandoc's default styles, with a Chinese font that exists on Windows."""
    from docx import Document
    from docx.oxml.ns import qn
    subprocess.run(["pandoc", "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(str(path))
    for style in doc.styles:
        try:
            rpr = style.element.get_or_add_rPr()
        except AttributeError:
            continue
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fonts)
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.save(str(path))


# ---------------------------------------------------------------------------
# Books
# ---------------------------------------------------------------------------
def load_book(n):
    folder = CONTENT / f"book{n}"
    meta = yaml.safe_load((folder / "book.yaml").read_text(encoding="utf-8"))
    days = [parse_day(f) for f in sorted(folder.glob("day*.txt"))]
    FORM_DAY.clear()
    for d in days:
        for v in d["vocab"]:
            FORM_DAY.setdefault(v["form"].lower(), d["day"])
            FORM_DAY.setdefault(v["head"].lower(), d["day"])
    for d in days:
        d["week_name"] = meta["weeks"][d["week"]]
    index = track_vocab(days)
    return meta, days, index


def book_stats(days, index):
    return {
        "days": len(days),
        "words": sum(1 for v in index if v["pos"] != "phr."),
        "phrases": sum(1 for v in index if v["pos"] == "phr."),
        "diffs": sum(len(d["diffs"]) for d in days),
        "core": sum(len(d["core"]) for d in days),
        "recurring": sum(1 for v in index if v["recur"]),
        "recur3": sum(1 for v in index if len(v["recur"]) >= 2),
        "story_words": sum(d["story_words"] for d in days),
    }


def check_book(n):
    meta, days, index = load_book(n)
    problems, warnings = validate(days)
    for w in warnings:
        print("warn:", w)
    for p in problems:
        print("ERROR:", p)
    s = book_stats(days, index)
    print(f"book{n}: {s['days']} days, {s['words']} words + {s['phrases']} phrases, "
          f"{s['diffs']} differences, {s['core']} core expressions, "
          f"{s['recurring']} items recur on later days")
    for d in days:
        print(f"  D{d['day']:02d} {d['story_words']:>3} story words, {d['word_count']:>2} words, "
              f"{d['phrase_count']} phrases, {len(d['diffs'])} diffs")
    if problems:
        sys.exit(1)
    return meta, days, index


def fill_counts(obj, stats):
    """Replace {WORDS} / {PHRASES} / {TOTAL} in book.yaml text with the real counts."""
    if isinstance(obj, str):
        return (obj.replace("{WORDS}", str(stats["words"])).replace("{PHRASES}", str(stats["phrases"]))
                .replace("{TOTAL}", str(stats["words"] + stats["phrases"])))
    if isinstance(obj, list):
        return [fill_counts(x, stats) for x in obj]
    if isinstance(obj, dict):
        return {k: fill_counts(v, stats) for k, v in obj.items()}
    return obj


def build_book(n):
    roadmap = load_syllabus()[0]["books"]  # loads Book 1 too, so do it before this book's own load
    meta, days, index = check_book(n)
    stats = book_stats(days, index)
    meta = fill_counts(meta, stats)
    ctx = {"book": meta, "days": days, "index": index, "stats": stats, "roadmap": roadmap, "day_titles": {d["day"]: d["title_cn"] for d in days},
           "alpha": sorted(index, key=lambda v: re.sub(r"[^a-z ]", "", v["head"].lower()) or v["head"]),
           "review_plan": REVIEW_PLAN, "regular_plan": REGULAR_PLAN}
    build_pdf("book.html.j2", ctx, meta["pdf_name"], meta["footer"],
              [("INDEX", r"Vocabulary Coverage"), ("RINDEX", r"Review Index"), ("LOG", r"Progress Log"),
               ("HANDOFF", r"Handoff to Book"), ("PROMPT", r"Prompt for Writing Book")])
    build_editable(ctx, meta["editable_name"])


def load_syllabus():
    """The 180-day roadmap; Book 1's vocabulary figures come from the finished lessons."""
    s = yaml.safe_load((CONTENT / "syllabus.yaml").read_text(encoding="utf-8"))
    meta, days, index = load_book(1)
    stats = book_stats(days, index)
    for b in s["books"]:
        b["words"] = b["words"].replace("{B1}", f"{stats['words']} 词 + {stats['phrases']} 短语")
    return s, meta, days, stats


def build_syllabus():
    s, meta, days, stats = load_syllabus()
    ctx = {"s": s, "b1": meta, "b1days": days, "b1stats": stats,
           "regular_plan": REGULAR_PLAN, "review_plan": REVIEW_PLAN}
    build_pdf("syllabus.html.j2", ctx, s["pdf_name"], s["footer"],
              [("LETTER", r"A Letter from Your Teacher"), ("WHO", r"Who It.s For"), ("HOW", r"How the Course Works"),
               ("MAP", r"The Six Books"), ("BOOKS", r"Book by Book"), ("B1", r"Book 1 Day by Day"),
               ("POINTS", r"Where Each Point Is Taught"), ("DIFF", r"Ten Big Differences"),
               ("TEST", r"Stage Assessments"), ("IPA", r"American English Sounds"),
               ("RES", r"Resources . Verification"), ("TRACK", r"Progress Tracker")])


def main():
    args = [a for a in sys.argv[1:]]
    books = sorted(int(p.name[4:]) for p in CONTENT.glob("book*") if p.is_dir())
    if args == ["check"]:
        for n in books:
            check_book(n)
        return
    ensure_fonts()
    targets = args or ["syllabus"] + [f"book{n}" for n in books]
    if "syllabus" in targets:
        build_syllabus()
    for n in books:
        if f"book{n}" in targets:
            build_book(n)


if __name__ == "__main__":
    try:
        main()
    except LessonError as e:
        sys.exit(f"ERROR: {e}")
